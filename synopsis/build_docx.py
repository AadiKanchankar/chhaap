"""Fill a copy of the PCCOE synopsis template with synopsis/synopsis.md.
Usage: .venv/bin/python synopsis/build_docx.py .   (writes synopsis/Chhaap_Synopsis.docx)"""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

ROOT = Path(sys.argv[1])
SYN = ROOT / "synopsis"
doc = docx.Document(SYN / "notice synopsis template_STUDENT.docx")
body = doc.element.body
P = lambda: [docx.text.paragraph.Paragraph(e, doc) for e in body.iterchildren(qn("w:p"))]

# ---- markdown sections ----
md = (SYN / "synopsis.md").read_text()
title = md.splitlines()[0].lstrip("% ").strip()
sec = {}
for block in re.split(r"^# ", md, flags=re.M)[1:]:
    head, _, rest = block.partition("\n")
    sec[head.strip()] = [b.strip() for b in rest.strip().split("\n\n") if b.strip()]

# ---- cover ----
paras = P()
label = next(p for p in paras if p.text.strip() == "Cover Page Template")._p
body.remove(label.getnext())  # empty paragraph under the label
body.remove(label)
for _ in range(3):  # make room for four member rows on one page
    body.remove(doc.tables[0]._tbl.getnext())
cover_title = next(p for p in paras if p.text.strip().startswith("“"))
for child in list(cover_title._p):
    if child.tag != qn("w:pPr"):
        cover_title._p.remove(child)
cover_title.paragraph_format.left_indent = 0
cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover_title.add_run(title)

t_members, t_sign = doc.tables[0], doc.tables[1]
members = [("123B1B166", "Keshav Shetshringi"), ("123B1B162", "Aadi Kanchankar"),
           ("123B1B168", "Swapnil Khillare"), ("123B1B169", "Pratham Kolettiwar")]
proto_row = t_members.rows[0]._tr
for prn, name in members:
    tr = copy.deepcopy(proto_row)
    t_members._tbl.append(tr)
    row = t_members.rows[-1]
    for cell, val in zip(row.cells, (prn, name)):
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        for r in p.runs:
            r._r.getparent().remove(r._r)
        run = p.add_run(val)
        run.bold = False
        run.font.size = Pt(12)
for tc in t_members._tbl.iter(qn("w:tcW")):
    if tc.get(qn("w:w")) == "2753":
        tc.set(qn("w:w"), "4200")
t_members._tbl.tblGrid[1].set(qn("w:w"), "4200")
for cell, who in ((t_sign.rows[0].cells[0], "Prof. Pallavi Nikumbh (Dhade)"), (t_sign.rows[0].cells[-1], "Dr. Anuradha Thakare")):
    gp = copy.deepcopy(cell.paragraphs[0]._p)
    cell._tc.append(gp)
    gpara = cell.paragraphs[-1]
    for r in gpara.runs:
        r._r.getparent().remove(r._r)
    gr = gpara.add_run(who)
    gr.bold = False
    gr.font.size = Pt(12)
for p in doc.paragraphs:  # cover: academic year
    if p.text.strip().startswith("ACADEMIC YEAR"):
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = "ACADEMIC YEAR 2026-27"
gid = next(p for p in doc.paragraphs if "Project Group ID" in p.text)
for r in gid.runs:
    if "-" in r.text:
        r.text = r.text.replace("------", " 62").replace("-", "")

# footer of section 2 carries the syllabus booklet line from the notice; keep only the page number
for t in doc.sections[1].footer._element.iter(qn("w:t")):
    if "Syllabus" in t.text or t.text == "27":
        t.text = ""

# ---- remove the instruction page, keep prototypes ----
paras = P()
start = next(i for i, p in enumerate(paras) if p.text.strip() == "Synopsis Template")
proto_main = copy.deepcopy(next(p for p in paras if p.text == "ABSTRACT")._p)
proto_sub = copy.deepcopy(next(p for p in paras if p.text == "Problem Statement")._p)
proto_sumlit = copy.deepcopy(next(p for p in paras if p.text.startswith("Summary of Literature"))._p)
proto_page_head = copy.deepcopy(paras[start]._p)
lit_tbl = copy.deepcopy(doc.tables[2]._tbl)
first = paras[start]._p
cover_break = copy.deepcopy(next(e for e in body.iterchildren(qn("w:p")) if e.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None))
for el in list(first.itersiblings()):
    if el.tag != qn("w:sectPr"):
        body.remove(el)
anchor = body[-1]  # final sectPr
body.remove(first)


def put(el):
    anchor.addprevious(el)
    return el


def set_text(p_el, text, bold=None):
    p = docx.text.paragraph.Paragraph(p_el, doc)
    rpr = p.runs[0]._r.find(qn("w:rPr")) if p.runs else None
    for r in p.runs:
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    if rpr is not None:
        run._r.insert(0, copy.deepcopy(rpr))
    if bold is not None:
        run.bold = bold
    return p


def main(text):
    return set_text(put(copy.deepcopy(proto_main)), text.upper())


def sub(text):
    return set_text(put(copy.deepcopy(proto_sub)), text)


INLINE = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")


def para(text, size=12, indent=0.5, hanging=0.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    text = INLINE.sub(lambda m: m.group(2), text)
    p = doc.add_paragraph(style="Body Text")
    put(p._p)
    p.alignment = align
    f = p.paragraph_format
    f.left_indent, f.right_indent = Inches(indent + hanging), Inches(0.3)
    f.first_line_indent = Inches(-hanging) if hanging else None
    f.space_after, f.line_spacing = Pt(6), 1.15
    p.add_run(text).font.size = Pt(size)
    return p


def blocks(name, skip=0):
    for b in sec[name][skip:]:
        if b.startswith("|"):
            table([[c.strip() for c in l.strip("|").split("|")] for l in b.splitlines() if not set(l) <= set("|:- ")])
        elif b.startswith("!["):
            figure()
        elif re.match(r"\d+\. ", b):
            for line in b.splitlines():
                para(line, hanging=0.25)
        else:
            para(b)


def fill_cell(cell, text, bold=False, size=11):
    p = cell.paragraphs[0]
    for r in p.runs:
        r._r.getparent().remove(r._r)
    p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(0.06)
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)


def table(rows, widths=None):
    ncol = len(rows[0])
    widths = widths or [9600 // ncol] * ncol
    t = doc.add_table(rows=len(rows), cols=ncol)
    put(t._tbl)
    t._tbl.remove(t._tbl.tblPr)
    t._tbl.insert(0, copy.deepcopy(lit_tbl.tblPr))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Twips(widths[j])
            fill_cell(c, val, bold=(i == 0))
    for g, w in zip(t._tbl.tblGrid.iterchildren(), widths):
        g.set(qn("w:w"), str(w))
    para("", size=6)
    return t


def figure():
    p = doc.add_paragraph(style="Body Text")
    put(p._p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SYN / "architecture.png"), width=Inches(6.5))
    para("Fig. 1. Chhaap architecture", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)


# ---- body ----
heads = ["Index", "Abstract", "Introduction", "Background", "Literature survey", "Methodology",
         "Expected conclusions", "Ethics and IP", "Mapping to course units", "Timeline", "References"]
if not any(e.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None for e in body.iterchildren(qn("w:p"))):
    put(cover_break)  # cover keeps its own section (no page footer)
while not "".join(cover_break.getprevious().itertext()).strip():  # trailing blanks spill onto an empty page
    body.remove(cover_break.getprevious())
set_text(put(proto_page_head), "Synopsis")
main("Index")
for i, h in enumerate(heads[1:], 2):
    para(f"{i}. {h}", indent=0.75)
main("Abstract"); blocks("Abstract")
prob = sec["Problem statement"]
main("Introduction"); para(prob[0])
main("Background")
sub("Problem Statement"); para(prob[1])
sub("Motivation/Need of project"); para(prob[2])
sub("Objectives"); blocks("Objectives")
sub("Scope of the Project"); blocks("Scope")
main("Literature survey"); blocks("Literature survey")
put(proto_sumlit)
lit = [
    ("High-resolution image synthesis with latent diffusion models [1]", "2022", "IEEE/CVF CVPR",
     "Denoising in the latent space of a pretrained autoencoder", "Lower training and sampling cost than pixel-space diffusion"),
    ("LoRA: Low-rank adaptation of large language models [2]", "2022", "ICLR",
     "Freezes pretrained weights and trains low-rank update matrices inserted into the layers", "Fewer trainable parameters"),
    ("Adding conditional control to text-to-image diffusion models [3]", "2023", "IEEE/CVF ICCV",
     "Adds spatial conditioning to a frozen diffusion model", "Inputs such as edges, depth or pose condition the output"),
    ("DreamBooth: Fine tuning text-to-image diffusion models for subject-driven generation [4]", "2023", "IEEE/CVF CVPR",
     "Fine-tunes a text-to-image model on a few images of a subject bound to a rare identifier", "Subject-driven generation from a few images"),
    ("LCM-LoRA: A universal Stable-Diffusion acceleration module [5]", "2023", "arXiv",
     "LoRA module distilled from a latent consistency model", "Attaches to Stable Diffusion variants without further training and permits few-step sampling"),
    ("GANs trained by a two time-scale update rule converge to a local Nash equilibrium [6]", "2017", "NeurIPS",
     "FID: compares Inception feature statistics of generated and real images", "Two time-scale update rule converges to a local Nash equilibrium"),
    ("Demystifying MMD GANs [7]", "2018", "ICLR",
     "KID, introduced with MMD GANs", "Unbiased alternative to FID that suits small samples"),
    ("Learning transferable visual models from natural language supervision [8]", "2021", "ICML (PMLR 139)",
     "CLIP: embeds images and text in one space", "Visual models learned from natural language supervision transfer to other tasks"),
    ("BLIP: Bootstrapping language-image pre-training for unified vision-language understanding and generation [9]", "2022", "ICML (PMLR 162)",
     "Bootstrapped language-image pre-training", "One model for vision-language understanding and generation; used here only for captions"),
    ("Potrace: a polygon-based tracing algorithm [10]", "2003", "P. Selinger (potrace.sourceforge.net)",
     "Polygon-based tracing", "Converts a bitmap into smooth Bezier outlines"),
    ("Learning texture manifolds with the periodic spatial GAN [11]", "2017", "ICML",
     "GAN that learns periodic textures", "Outputs of arbitrary size"),
    ("Tiled Diffusion [12]", "2025", "IEEE/CVF CVPR",
     "Diffusion that generates tiles that connect", "Self-tiling, many-to-many tile sets and 360-degree images; no treatment of motif meaning"),
]
t = docx.table.Table(put(lit_tbl), doc)
proto = t.rows[0]._tr
for i, row in enumerate(lit, 1):
    tr = copy.deepcopy(proto)
    t._tbl.append(tr)
    cells = t.rows[-1].cells
    for c in cells:  # template header cells hold 2 paragraphs in places
        for extra in c.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
    for c, val in zip(cells, (str(i),) + row):
        fill_cell(c, val, size=10)
lit_w = [700, 2400, 1300, 1250, 2000, 1950]  # template's 1260-twip title column wraps every word
for g, w in zip(t._tbl.tblGrid.iterchildren(), lit_w):
    g.set(qn("w:w"), str(w))
for row in t.rows:
    for c, w in zip(row.cells, lit_w):
        c.width = Twips(w)
for tr in t._tbl.iter(qn("w:trPr")):
    for h in tr.iter(qn("w:trHeight")):
        h.set(qn("w:val"), "300")
para("", size=6)

main("Methodology")
ps = sec["Proposed system"]
sub("Block diagram"); figure(); para(ps[0]); para(ps[1])
sub("Hardware and Software requirement"); blocks("Hardware and software")
sub("Working principle"); para(ps[3]); blocks("Methodology")
sub("Dataset"); blocks("Dataset")
sub("Evaluation plan"); blocks("Evaluation plan")
main("Expected conclusions"); blocks("Expected outcomes")
sub("Future work"); blocks("Future work")
main("Ethics and IP"); blocks("Ethics and IP")
main("Mapping to course units"); blocks("Mapping to course units")
main("Timeline"); blocks("Timeline")
main("References")
for b in sec["References"]:
    para(b, size=11, hanging=0.4, align=WD_ALIGN_PARAGRAPH.LEFT)

assert ps[2].startswith("![")
used = {"Abstract", "Problem statement", "Objectives", "Scope", "Literature survey", "Proposed system", "Hardware and software",
        "Methodology", "Dataset", "Evaluation plan", "Expected outcomes", "Future work", "Ethics and IP",
        "Mapping to course units", "Timeline", "References"}
assert used == set(sec), set(sec) ^ used
doc.save(SYN / "Chhaap_Synopsis.docx")
print("saved")
