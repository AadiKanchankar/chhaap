"""Convert paper/paper.tex to paper/Chhaap_Paper.docx (opens directly in Google Docs).
Usage: .venv/bin/python paper/build_docx.py
paper.tex is never modified: IEEEtran-only macros are rewritten in a temp copy first."""
import copy
import math
import re
import subprocess
import tempfile
from pathlib import Path

import docx
import pypandoc
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PAPER = Path(__file__).resolve().parent
OUT = PAPER / "Chhaap_Paper.docx"


AUTHOR_RE = re.compile(r"\\author\{(.*?)\}\s*\\maketitle", re.S)


def author_rows(tex: str):
    """\\author{...} -> rows (split on \\linebreakand) of (name, [affiliation lines])."""
    body = AUTHOR_RE.search(tex).group(1)
    return [[(n, [ln.strip().replace("\\&", "&") for ln in a.split("\\\\")])
             for n, a in re.findall(r"\\IEEEauthorblockN\{([^{}]*)\}\s*\\IEEEauthorblockA\{([^{}]*)\}", row)]
            for row in body.split("\\linebreakand")]


def preprocess(tex: str) -> str:
    tex = re.sub(r"\\makeatletter.*?\\makeatother\n?", "", tex, flags=re.S)
    # authors are added after pandoc as a table (author_table); pandoc would stack them in one column
    tex = AUTHOR_RE.sub(r"\\maketitle", tex)
    # \mathrm{SDR} becomes per-letter runs in OMML (spaced italic in Word); \text{} is one upright run
    tex = tex.replace("\\mathrm{", "\\text{")
    tex = re.sub(r"\\begin\{IEEEkeywords\}\s*(.*?)\s*\\end\{IEEEkeywords\}",
                 r"\\textbf{\\textit{Index Terms}}---\1", tex, flags=re.S)
    # unwrap \resizebox{..}{..}{ <tabular> }
    tex = re.sub(r"\\resizebox\{[^{}]*\}\{[^{}]*\}\{%?\s*(\\begin\{tabular\}.*?\\end\{tabular\})\s*\}",
                 r"\1", tex, flags=re.S)
    # bibliography comes from --citeproc; drop the BibTeX commands
    return re.sub(r"\\bibliography(style)?\{[^}]*\}\n?", "", tex)


def reference_docx(path: Path):
    """Pandoc's default styles, restyled to look like a paper (Times, black headings)."""
    subprocess.run([pypandoc.get_pandoc_path(), "-o", str(path), "--print-default-data-file",
                    "reference.docx"], check=True)
    d = docx.Document(path)
    sizes = {"Title": 20, "Heading 1": 12, "Heading 2": 11, "Heading 3": 10}
    for st in d.styles:
        if st.type != 1:  # paragraph styles only
            continue
        st.font.name = "Times New Roman"
        rpr = st.element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), "Times New Roman")
        for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):  # theme fonts override w:ascii
            rf.attrib.pop(qn("w:" + a), None)
        st.font.color.rgb = RGBColor(0, 0, 0)
        if st.name in sizes:
            st.font.size = Pt(sizes[st.name])
            st.font.bold = True
            st.font.italic = st.name == "Heading 2"
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Bibliography"):
        if name in [s.name for s in d.styles]:
            d.styles[name].font.size = Pt(10)
    for name in ("Title", "Author", "Abstract Title", "Image Caption", "Table Caption", "Captioned Figure"):
        if name in [s.name for s in d.styles]:
            d.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.save(path)


def author_table(path: Path, rows):
    """Insert a borderless, centred author grid after the title. Rows of different lengths share a
    grid of lcm columns (3 + 2 -> 6 columns: spans of 2 and 3)."""
    d = docx.Document(path)
    grid = math.lcm(*map(len, rows))
    sec = d.sections[0]
    # pandoc's sectPr carries no page size; fall back to Word's Letter with 1-inch margins
    width = (sec.page_width - sec.left_margin - sec.right_margin) if sec.page_width else Inches(6.5)
    t = d.add_table(rows=len(rows), cols=grid)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = t._tbl.tblPr
    for el in tbl_pr.findall(qn("w:tblStyle")):
        tbl_pr.remove(el)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "nil")
        borders.append(b)
    tbl_pr.append(borders)
    for c in t.columns:
        c.width = width // grid
    for r, authors in zip(t.rows, rows):
        span = grid // len(authors)
        for i, (name, lines) in enumerate(authors):
            cell = r.cells[i * span].merge(r.cells[(i + 1) * span - 1]) if span > 1 else r.cells[i * span]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(name)
            run.bold = True
            run.font.size = Pt(10)
            for ln in lines:
                run.add_break()
                run = p.add_run(ln)
                run.font.size = Pt(9)
    title = next(p for p in d.paragraphs if p.style.name == "Title")
    title._p.addnext(t._tbl)
    d.save(path)


def two_column(path: Path):
    """IEEE look: Letter page, title block in one column, body in two; images and tables fit a column."""
    d = docx.Document(path)
    for sec in d.sections:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(0.625)
        sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1)
    col = Inches((8.5 - 2 * 0.625 - 0.25) / 2)
    for st in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if st in [x.name for x in d.styles]:
            d.styles[st].font.size, d.styles[st].font.name = Pt(10), "Times New Roman"
    body = d.element.body
    last = body.find(qn("w:sectPr"))
    keys = next(p for p in d.paragraphs if p.text.startswith("Index Terms"))
    head = copy.deepcopy(last)  # section that ends after the keywords: one column
    for c in head.findall(qn("w:cols")):
        head.remove(c)
    cols1 = OxmlElement("w:cols"); cols1.set(qn("w:num"), "1"); head.append(cols1)
    keys._p.get_or_add_pPr().append(head)
    for c in last.findall(qn("w:cols")):
        last.remove(c)
    cols2 = OxmlElement("w:cols"); cols2.set(qn("w:num"), "2"); cols2.set(qn("w:space"), str(int(0.25 * 1440)))
    last.append(cols2)
    t = OxmlElement("w:type"); t.set(qn("w:val"), "continuous"); last.insert(0, t)
    for shape in d.inline_shapes:  # every figure fits one column, aspect kept
        if shape.width > col:
            shape.height, shape.width = int(shape.height * col / shape.width), col
    for tbl in d.tables[1:]:  # the author grid (first table) spans the page
        w = tbl._tbl.tblPr.find(qn("w:tblW"))
        if w is None:
            w = OxmlElement("w:tblW"); tbl._tbl.tblPr.append(w)
        w.set(qn("w:type"), "pct"); w.set(qn("w:w"), "5000")
        for cell in (c for r in tbl.rows for c in r.cells):
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(8)
    d.save(path)


def main():
    raw = (PAPER / "paper.tex").read_text()
    tex = preprocess(raw)
    # fails if paper.tex gains a macro shape preprocess() does not rewrite
    assert not re.search(r"\\IEEE|\\resizebox|\\linebreakand", tex), "unhandled IEEEtran macro"
    with tempfile.TemporaryDirectory() as tmp:
        src, ref = Path(tmp) / "paper.tex", Path(tmp) / "reference.docx"
        src.write_text(tex)
        reference_docx(ref)
        pypandoc.convert_file(
            str(src), "docx+native_numbering", format="latex", outputfile=str(OUT),
            extra_args=["--citeproc", f"--bibliography={PAPER / 'refs.bib'}",
                        f"--csl={PAPER / 'ieee-min.csl'}", "--number-sections",
                        f"--resource-path={PAPER}", f"--reference-doc={ref}",
                        "--metadata=reference-section-title:References"],
            cworkdir=str(PAPER))
    author_table(OUT, author_rows(raw))
    two_column(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
